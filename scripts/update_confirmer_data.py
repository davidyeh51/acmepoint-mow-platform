# -*- coding: utf-8 -*-
import openpyxl
import json
import sys

if hasattr(sys.stdout, 'reconfigure'):
    sys.stdout.reconfigure(encoding='utf-8')

# 1. 更新 Excel
excel_path = '維運管理系統_案場資料庫_完整資料表.xlsx'
wb = openpyxl.load_workbook(excel_path)
ws = wb['客戶服務_VOC']

headers = [cell.value for cell in ws[1]]
print("Excel headers:", headers)

if '確認者' not in headers:
    col_confirmer = len(headers) + 1
    col_time = len(headers) + 2
    ws.cell(row=1, column=col_confirmer, value='確認者')
    ws.cell(row=1, column=col_time, value='確認時間')
    
    for r in range(2, ws.max_row + 1):
        status = ws.cell(row=r, column=12).value
        handler = ws.cell(row=r, column=7).value
        close_date = ws.cell(row=r, column=10).value
        if status == '結案' and handler:
            ws.cell(row=r, column=col_confirmer, value=f"{handler}(結案審核)")
            ws.cell(row=r, column=col_time, value=str(close_date) if close_date else '')
        else:
            ws.cell(row=r, column=col_confirmer, value='')
            ws.cell(row=r, column=col_time, value='')
            
    wb.save(excel_path)
    print("Excel VOC sheet updated with 確認者 column successfully.")
else:
    print("Excel already has 確認者 column.")

# 2. 更新 om_platform_data.js
with open('om_platform_data.js', 'r', encoding='utf-8') as f:
    content = f.read()

prefix = 'window.OM_DATA = '
if content.startswith(prefix):
    json_str = content[len(prefix):].rstrip(';\n')
    data = json.loads(json_str)
    
    for item in data.get('voc_list', []):
        if 'confirmer' not in item or not item['confirmer']:
            if item.get('status') == '結案' and item.get('handler'):
                item['confirmer'] = f"{item['handler']}(結案審核)"
                item['confirmed_at'] = item.get('close_date', '')
            else:
                item['confirmer'] = ''
                item['confirmed_at'] = ''
                
    with open('om_platform_data.js', 'w', encoding='utf-8') as f_out:
        f_out.write(prefix + json.dumps(data, ensure_ascii=False) + ';')
    print("om_platform_data.js updated with confirmer fields.")
